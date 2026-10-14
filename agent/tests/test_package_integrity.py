"""Every workpaper this produces must open in Excel.

A real output was refused by Excel with "cannot be opened or repaired" --
a message that names nothing. Three separate causes were found behind it,
each now fixed and each tested here, plus a final check on every package
that refuses to ship anything it can see is broken.
"""

from __future__ import annotations

import re
import shutil
import zipfile
from pathlib import Path

import openpyxl

from agent.build_cy_workpaper import build_cy_workpaper
from agent.ooxml import ooxml_utils as ox
from agent.tests.test_build_cy_workpaper import (
    _manifests, _real_pdf, _results, _spec, _two_manifests, _two_sample_results,
)


def _problems(xlsx: Path, tmp: Path) -> list[str]:
    work = tmp / ("chk_" + xlsx.stem)
    ox.extract_template(str(xlsx), str(work))
    return ox.package_problems(str(work))


def _py_like_the_second_control(path: Path, *, printer_settings: bool = True) -> Path:
    """Tabs '1' and 'IA Lead', no pictures anywhere (so no .png content
    type), and -- like most Excel-saved tabs -- printer settings on the
    sample tab, referenced by r:id from its <pageSetup>."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "1"
    ws["A1"] = "Invoice Number"
    ws["A2"] = "PY-111"
    for i in range(20):
        ws.cell(50 + i, 1, i).number_format = f"0.{'0' * (i % 8)}#"
    wb.create_sheet("IA Lead")["A1"] = "Lead"
    wb.save(path)
    if not printer_settings:
        return path

    work = path.parent / "ps_work"
    ox.extract_template(str(path), str(work))
    num = ox.resolve_sheets(str(work))["1"]
    (work / "xl" / "printerSettings").mkdir()
    (work / "xl" / "printerSettings" / "printerSettings1.bin").write_bytes(b"\x00" * 64)
    rels = work / "xl" / "worksheets" / "_rels" / f"sheet{num}.xml.rels"
    rels.parent.mkdir(exist_ok=True)
    rels.write_text(
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rIdPS" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
        'relationships/printerSettings" Target="../printerSettings/printerSettings1.bin"/>'
        "</Relationships>"
    )
    sheet = work / "xl" / "worksheets" / f"sheet{num}.xml"
    xml = sheet.read_text()
    if "xmlns:r=" not in xml.split(">", 1)[0]:
        xml = xml.replace("<worksheet ", '<worksheet xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" ', 1)
    xml = re.sub(r"(<pageMargins\b[^>]*/>)", r'\1<pageSetup orientation="landscape" r:id="rIdPS"/>', xml, count=1)
    sheet.write_text(xml)
    ct = (work / "[Content_Types].xml").read_text()
    ct = ct.replace("</Types>", '<Default Extension="bin" ContentType="application/vnd.openxmlformats-'
                                 'officedocument.spreadsheetml.printerSettings"/></Types>', 1)
    (work / "[Content_Types].xml").write_text(ct)
    ox.repackage(str(work), str(path))
    shutil.rmtree(work)
    return path


def test_the_second_controls_shape_with_two_samples_opens(tmp_path: Path):
    # Two selections means the PY tab is COPIED. The copy carried the
    # printer-settings r:id without the relationship it names.
    _py_like_the_second_control(tmp_path / "PY.xlsx")
    _real_pdf(tmp_path / "invoice.pdf")
    out = build_cy_workpaper(_spec(), _two_sample_results(), "PY.xlsx", tmp_path / "o",
                             support_dir=tmp_path, sample_manifests=_two_manifests())
    assert out.mirrored_py, out.fallback_reason
    assert _problems(out.path, tmp_path) == []
    # The copied tab keeps its print setup, with a part of its own.
    z = zipfile.ZipFile(out.path)
    assert len([n for n in z.namelist() if n.startswith("xl/printerSettings/")]) == 2
    wb = openpyxl.load_workbook(out.path)
    assert wb.sheetnames == ["Summary", "1", "2", "IA Lead"]


def test_a_py_with_no_pictures_still_gets_an_image_content_type(tmp_path: Path):
    _py_like_the_second_control(tmp_path / "PY.xlsx", printer_settings=False)
    _real_pdf(tmp_path / "invoice.pdf")
    out = build_cy_workpaper(_spec(), _results(), "PY.xlsx", tmp_path / "o",
                             support_dir=tmp_path, sample_manifests=_manifests())
    assert out.mirrored_py, out.fallback_reason
    ct = zipfile.ZipFile(out.path).read("[Content_Types].xml").decode()
    assert re.search(r'<Default[^>]*Extension="png"', ct)


def test_a_drawing_declared_the_non_excel_way_is_reused_not_duplicated(tmp_path: Path):
    # <drawing xmlns:r="..." r:id="rId1"/> was missed, so a second <drawing>
    # with an undeclared prefix was added: malformed XML.
    from agent.tests.test_layout import _MEASURED, _py_with_pictures

    py = _py_with_pictures(tmp_path, _MEASURED[:2])
    _real_pdf(tmp_path / "invoice.pdf")
    out = build_cy_workpaper(_spec(), _results(), py.name, tmp_path / "o",
                             support_dir=tmp_path, sample_manifests=_manifests())
    assert out.mirrored_py, out.fallback_reason
    z = zipfile.ZipFile(out.path)
    for n in z.namelist():
        if re.fullmatch(r"xl/worksheets/sheet\d+\.xml", n):
            assert z.read(n).decode().count("<drawing") <= 1, n
    assert _problems(out.path, tmp_path) == []


def test_a_file_that_fails_the_final_check_falls_back_instead_of_shipping(tmp_path: Path, monkeypatch):
    from agent.tests.test_build_cy_workpaper import _template

    _template(tmp_path / "PY.xlsx")
    _real_pdf(tmp_path / "invoice.pdf")
    monkeypatch.setattr(ox, "package_problems", lambda _w: ["sheet9.xml refers to rId7, which is not defined"])
    out = build_cy_workpaper(_spec(), _results(), "PY.xlsx", tmp_path / "o",
                             support_dir=tmp_path, sample_manifests=_manifests())
    assert not out.mirrored_py
    assert out.path.exists()
    assert "integrity check" in out.fallback_reason and "rId7" in out.fallback_reason


def test_the_check_names_the_problems_excel_would_not(tmp_path: Path):
    from agent.tests.test_build_cy_workpaper import _template

    path = _template(tmp_path / "PY.xlsx")
    work = tmp_path / "w"
    ox.extract_template(str(path), str(work))
    assert ox.package_problems(str(work)) == []
    sheet = next((work / "xl" / "worksheets").glob("sheet*.xml"))
    sheet.write_text(sheet.read_text().replace("</worksheet>", '<pageSetup r:id="rIdGone"/></worksheet>'))
    problems = ox.package_problems(str(work))
    assert problems and ("rIdGone" in " ".join(problems) or "well-formed" in " ".join(problems))
