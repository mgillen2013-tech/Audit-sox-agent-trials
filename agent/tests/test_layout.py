"""Exhibit placement on a Sample tab: mirrored from the PY, never stacked.

A real output stacked four exhibits at full page size in one column, the
fourth starting 29 inches down -- ~140 rows of scrolling -- while the PY for
the same control laid its exhibits out three across and two down.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
from PIL import Image

from agent.ooxml import ooxml_utils as ox
from agent.ooxml.layout import (
    EMU_PER_INCH, GAP, Slot, TemplateLayout, plan_sample_layout, read_template_layout,
)

PAGE = (1236, 1622)      # a scanned letter page, as rendered
WIDE = (1293, 695)       # a payment-screen screenshot


def _no_overlaps(boxes: list[Slot]) -> bool:
    return not any(a.overlaps(b) for i, a in enumerate(boxes) for b in boxes[i + 1:])


def _plan(layout: TemplateLayout, sizes, narrative_h=int(2.7 * EMU_PER_INCH)):
    return plan_sample_layout(
        layout, sizes,
        narrative_height_for_width=lambda _w: narrative_h,
        default_narrative_width=5_545_854,
    )


# --------------------------------------------------------------------------
# Planning (pure geometry)
# --------------------------------------------------------------------------


def test_without_py_slots_exhibits_go_side_by_side_not_down_one_column():
    narrative, boxes = _plan(TemplateLayout(), [PAGE, PAGE, PAGE, WIDE])
    deepest = max(b.bottom for b in boxes) / EMU_PER_INCH
    # The single-column stack put the last of these ~33in down.
    assert deepest < 12, deepest
    # The first two share the narrative's band, to its right.
    assert boxes[0].y == boxes[1].y == narrative.y
    assert boxes[0].x >= narrative.right
    assert _no_overlaps([narrative, *boxes])


def test_an_exhibit_under_a_short_narrative_sits_right_under_it():
    # Placing the next row under the tallest thing in the row left a blank
    # patch beneath the narrative.
    narrative, boxes = _plan(TemplateLayout(), [PAGE, PAGE, PAGE])
    third = boxes[2]
    assert third.x < narrative.right
    assert third.y == narrative.bottom + GAP


def test_exhibits_fill_the_py_slots_in_reading_order_and_fit_inside_them():
    slots = [Slot(6_000_000, 600_000, 4_000_000, 4_000_000),
             Slot(10_500_000, 600_000, 4_000_000, 4_000_000)]
    layout = TemplateLayout(narrative=Slot(0, 600_000, 5_500_000, 2_000_000), images=slots)
    _narrative, boxes = _plan(layout, [PAGE, WIDE], narrative_h=1_000_000)
    for box, slot in zip(boxes, slots):
        assert (box.x, box.y) == (slot.x, slot.y)
        assert box.cx <= slot.cx and box.cy <= slot.cy
    # Unstretched: the page keeps its shape.
    assert abs(boxes[0].cx / boxes[0].cy - PAGE[0] / PAGE[1]) < 0.01


def test_more_exhibits_than_py_slots_continue_below_without_overlapping():
    slots = [Slot(6_000_000, 600_000, 4_000_000, 5_000_000),
             Slot(10_500_000, 600_000, 4_000_000, 5_000_000)]
    layout = TemplateLayout(narrative=Slot(0, 600_000, 5_500_000, 2_000_000), images=slots)
    narrative, boxes = _plan(layout, [PAGE] * 5, narrative_h=1_000_000)
    assert len(boxes) == 5
    assert _no_overlaps([narrative, *boxes])
    # Extra exhibits use the PY's slot size, not full page size.
    assert all(b.cy <= 5_000_000 for b in boxes)


def test_text_that_outgrows_its_py_box_pushes_exhibits_down_rather_than_under_them():
    layout = TemplateLayout(
        narrative=Slot(0, 600_000, 3_000_000, 1_000_000),
        images=[Slot(0, 1_800_000, 3_000_000, 3_000_000)],
    )
    narrative, boxes = _plan(layout, [PAGE], narrative_h=4_000_000)
    assert narrative.cy == 4_000_000  # grown, never clipped
    assert not boxes[0].overlaps(narrative)


# --------------------------------------------------------------------------
# Reading the PY drawing
# --------------------------------------------------------------------------


def _py_with_pictures(tmp_path: Path, anchors: list[tuple[int, int, int, int]]) -> Path:
    """A PY-shaped workbook whose Sample tab has pictures at the given
    (col_from, row_from, col_to, row_to) cells, written as twoCellAnchor --
    the anchor type Excel itself writes."""
    page = tmp_path / "page.png"
    Image.new("RGB", PAGE, "white").save(page)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sample"
    ws["A1"] = "Invoice Number"
    for c0, r0, c1, r1 in anchors:
        img = XLImage(str(page))
        img.anchor = TwoCellAnchor(_from=AnchorMarker(col=c0, row=r0), to=AnchorMarker(col=c1, row=r1))
        ws.add_image(img)
    wb.create_sheet("Population")["A5"] = "Invoice Number"
    path = tmp_path / "PY.xlsx"
    wb.save(path)
    return path


# Measured from the real PY Sample tab (docs/layout_must_be_mirrored_not_computed.md)
_MEASURED = [(6, 3, 19, 24), (19, 2, 36, 40), (0, 26, 6, 70), (7, 25, 19, 68), (21, 40, 34, 68)]


def _read(tmp_path: Path, path: Path) -> TemplateLayout:
    work = str(tmp_path / "w")
    ox.extract_template(str(path), work)
    sheet = f"{work}/xl/worksheets/sheet{ox.resolve_sheets(work)['Sample']}.xml"
    return read_template_layout(work, sheet, ox.parse_column_widths(sheet))


def test_py_picture_positions_are_read_as_slots_in_reading_order(tmp_path: Path):
    layout = _read(tmp_path, _py_with_pictures(tmp_path, _MEASURED))
    assert layout.usable and len(layout.images) == 5
    # Top band left-to-right, then the next band left-to-right.
    xs_top = [s.x for s in layout.images[:2]]
    assert xs_top == sorted(xs_top)
    assert layout.images[0].y < layout.images[2].y
    # Cell anchors became real sizes: rows 3-24 at 15pt is about 4.4in.
    assert 4.0 < layout.images[0].cy / EMU_PER_INCH < 4.8


def test_a_previous_cy_output_is_not_mirrored(tmp_path: Path):
    # Its drawing is the old single-column stack; copying it would keep
    # exactly the layout being fixed.
    from agent.tests.test_build_cy_workpaper import (
        _manifests, _real_pdf, _results, _spec, _template,
    )
    from agent.build_cy_workpaper import build_cy_workpaper

    _template(tmp_path / "PY0.xlsx")
    _real_pdf(tmp_path / "invoice.pdf")
    out = build_cy_workpaper(_spec(), _results(), "PY0.xlsx", tmp_path / "o",
                             support_dir=tmp_path, sample_manifests=_manifests())
    layout = _read(tmp_path, out.path)
    assert layout.generated and not layout.usable


def test_an_unreadable_drawing_falls_back_instead_of_failing(tmp_path: Path):
    path = _py_with_pictures(tmp_path, _MEASURED[:1])
    work = str(tmp_path / "w")
    ox.extract_template(str(path), work)
    for f in Path(work, "xl", "drawings").glob("drawing*.xml"):
        f.write_text("<not xml")
    sheet = f"{work}/xl/worksheets/sheet{ox.resolve_sheets(work)['Sample']}.xml"
    assert not read_template_layout(work, sheet, ox.parse_column_widths(sheet)).usable


# --------------------------------------------------------------------------
# End to end
# --------------------------------------------------------------------------


def test_the_workpaper_puts_the_exhibit_where_the_py_had_its_picture(tmp_path: Path):
    import re
    import zipfile

    from agent.build_cy_workpaper import build_cy_workpaper
    from agent.tests.test_build_cy_workpaper import _manifests, _real_pdf, _results, _spec

    py = _py_with_pictures(tmp_path, _MEASURED)
    expected = _read(tmp_path, py).images[0]
    _real_pdf(tmp_path / "invoice.pdf")
    out = build_cy_workpaper(_spec(), _results(), py.name, tmp_path / "o",
                             support_dir=tmp_path, sample_manifests=_manifests())
    assert out.mirrored_py, out.fallback_reason
    z = zipfile.ZipFile(out.path)
    drawings = [z.read(n).decode() for n in z.namelist() if re.fullmatch(r"xl/drawings/drawing\d+\.xml", n)]
    placed = [
        (int(x), int(y))
        for d in drawings
        for x, y in re.findall(r'<xdr:pos x="(\d+)" y="(\d+)"/>\s*<xdr:ext[^>]*/>\s*<xdr:pic>', d)
    ]
    assert (expected.x, expected.y) in placed, (expected, placed)
