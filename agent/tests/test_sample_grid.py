"""The sampled item's row: under the PY's headings, wherever they are.

Two assumptions this replaces, both live in real outputs:
- headings on row 1, item on row 2, on every PY sample tab;
- values written into columns A, B, C... in the LISTING's order. A real
  156 output had every value from column F onward one column left of its
  heading -- the payment amount under "Void Date" -- because the item's
  blank Void Date was dropped and everything after it slid over.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl

from agent.ooxml.sample_grid import plan_grid

# The real 156 PY sample tab headings (row 1), as read from its workpaper.
PY_156 = [
    "Match Doc Ty F0413.DCTM", "Check/ Item F0413.DOCM", "Payee Number F0413.PYE",
    "Payee Number F0413.PYE", "Check/ Itm Date F0413.DMTJ", "Void Date F0413.VDGJ",
    "Total Payment Amount F0413.PAAP", "Do Ty F0414.DCT", "Document Number F0414.DOC",
    "Pay Itm F0414.SFX", "Invoice Payment Amount F0414.PAAP",
]


def _cells(rows: dict[int, list[str]]) -> dict[str, tuple]:
    out = {}
    for r, texts in rows.items():
        for i, t in enumerate(texts):
            if t:
                out[f"{chr(65 + i)}{r}"] = (None, t, False)
    return out


def _listing(names: list[str]) -> dict[str, str]:
    return {chr(65 + i): n for i, n in enumerate(names)}


def test_a_dropped_blank_column_no_longer_shifts_everything_after_it():
    # This year's item: Void Date blank (so absent), and the duplicate
    # "Payee Number" deduplicated by intake.
    listing = _listing([
        "Match Doc Ty F0413.DCTM", "Check/ Item F0413.DOCM", "Payee Number F0413.PYE",
        "Payee Number F0413.PYE (2)", "Check/ Itm Date F0413.DMTJ",
        "Total Payment Amount F0413.PAAP", "Do Ty F0414.DCT",
    ])
    grid = plan_grid(_cells({1: PY_156}), listing)
    assert grid.header_row == 1 and grid.data_row == 2
    # Total Payment Amount (listing F) lands under its own heading, G -- not
    # under F, "Void Date", where the old positional write put it.
    assert grid.column_for["F"] == "G"
    assert grid.column_for["G"] == "H"
    assert grid.column_for["D"] == "D"  # the second Payee Number
    assert not grid.headings_to_write


def test_a_column_the_py_has_no_heading_for_goes_where_the_py_had_it():
    # The vendor name: no PY heading of its own, but an unused PY column sits
    # between its neighbours -- use it, rather than the far right.
    listing = _listing([
        "Payee Number F0413.PYE", "Alpha Name", "Check/ Itm Date F0413.DMTJ",
    ])
    grid = plan_grid(_cells({1: PY_156}), listing)
    assert grid.column_for == {"A": "C", "B": "D", "C": "E"}


def test_a_column_with_nowhere_between_its_neighbours_is_added_with_its_heading():
    listing = _listing(["Payee Number F0413.PYE", "Approver Email", "Payee Number F0413.PYE (2)"])
    grid = plan_grid(_cells({1: PY_156}), listing)
    added = grid.column_for["B"]
    assert added not in [chr(65 + i) for i in range(len(PY_156))]
    assert grid.headings_to_write == {added: "Approver Email"}


def test_reordered_listing_columns_follow_their_headings():
    listing = _listing(["Do Ty F0414.DCT", "Match Doc Ty F0413.DCTM", "Payee Number F0413.PYE"])
    grid = plan_grid(_cells({1: PY_156}), listing)
    assert grid.column_for == {"A": "H", "B": "A", "C": "C"}


def test_headings_below_a_title_are_found_and_the_title_is_kept():
    cells = _cells({
        1: ["SS.XXX.01 Sample Testing"],
        2: ["Prepared by: J. Smith"],
        3: ["Invoice Number", "Vendor", "Amount"],
        4: ["PY-111", "Old Vendor", "100"],
    })
    grid = plan_grid(cells, _listing(["Invoice Number", "Vendor", "Amount"]))
    assert (grid.header_row, grid.data_row) == (3, 4)
    assert any("row 3" in n for n in grid.notes)


def test_no_matching_headings_writes_the_listings_own_names():
    # Values under someone else's labels look right and are wrong.
    grid = plan_grid(_cells({1: ["Foo", "Bar", "Baz"]}), _listing(["Invoice Number", "Amount"]))
    assert (grid.header_row, grid.data_row) == (1, 2)
    assert grid.headings_to_write == {"A": "Invoice Number", "B": "Amount"}


def test_duplicate_listing_headings_are_both_kept(tmp_path: Path):
    from agent.intake import read_excel_rows

    wb = openpyxl.Workbook()
    wb.active.append(["Payee Number", "Payee Number", "Amount"])
    wb.active.append(["881477", "Premiere Onboard LLC", 30000])
    wb.save(tmp_path / "listing.xlsx")
    row = read_excel_rows(tmp_path / "listing.xlsx")[0]
    assert row["Payee Number"] == "881477"
    assert row["Payee Number (2)"] == "Premiere Onboard LLC"


def test_end_to_end_a_py_with_a_title_above_its_headings(tmp_path: Path):
    from agent.build_cy_workpaper import build_cy_workpaper
    from agent.schemas import SampleItem, SamplePopulationManifest
    from agent.tests.test_build_cy_workpaper import _real_pdf, _results, _spec

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sample"
    ws["A1"] = "SS.XXX.01 Sample Testing"
    ws.append([])
    ws.append(["Vendor", "Invoice Number"])  # row 3; note: listing order differs
    ws.append(["Old Vendor", "PY-111"])
    for i in range(20):
        ws.cell(50 + i, 1, i).number_format = f"0.{'0' * (i % 8)}#"
    wb.create_sheet("Population")
    wb.save(tmp_path / "PY.xlsx")
    _real_pdf(tmp_path / "invoice.pdf")
    manifests = {"TS-1": SamplePopulationManifest(
        test_step_id="TS-1", sample_size=1,
        samples=[SampleItem(sample_id="1", test_step_id="TS-1", identifying_details="x",
                            key_fields={"Invoice Number": "21022452", "Vendor": "Acme"})],
    )}
    out = build_cy_workpaper(_spec(), _results(), "PY.xlsx", tmp_path / "o",
                             support_dir=tmp_path, sample_manifests=manifests)
    assert out.mirrored_py, out.fallback_reason
    got = openpyxl.load_workbook(out.path)["Sample"]
    assert got["A1"].value == "SS.XXX.01 Sample Testing"   # title kept
    assert (got["A3"].value, got["B3"].value) == ("Vendor", "Invoice Number")  # headings kept
    assert (got["A4"].value, got["B4"].value) == ("Acme", "21022452")  # by name, on row 4
    assert any("row 3" in w for w in out.warnings)
    # The red box over the invoice number moved with it: column B, row 4.
    import re
    import zipfile

    z = zipfile.ZipFile(out.path)
    boxes = [
        tuple(map(int, m.groups()))
        for n in z.namelist() if re.fullmatch(r"xl/drawings/drawing\d+\.xml", n)
        for m in re.finditer(
            r'<xdr:pos x="(\d+)" y="(\d+)"/>\s*<xdr:ext cx="\d+" cy="\d+"/>\s*<xdr:sp[^>]*>\s*'
            r'<xdr:nvSpPr>\s*<xdr:cNvPr id="\d+" name="Rectangle raw', z.read(n).decode())
    ]
    assert boxes == [(64 * 9525, 3 * 15 * 12700)]  # col B left edge, row 4 top
