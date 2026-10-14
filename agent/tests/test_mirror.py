"""The mirror engine: a CY workpaper made by editing a copy of the PY.

The fixture PY is built to the grammar shared by every real PY seen so
far (CI.9, MAR.17, AP.156): a lead tab with one description box, a tab per
selection with headings and a data row, exhibit slots with red frames,
tickmark label + box pairs, and one conclusion box per test step in the
house style ("Test Step 1: ... / Test Satisfied."). Real PYs are never
committed; this reproduces their structure, not their content.
"""

from __future__ import annotations

import json
import re
import zipfile
from pathlib import Path

import openpyxl
import pytest
from openpyxl.drawing.image import Image as XLImage
from PIL import Image, ImageDraw, ImageFont

from agent.mirror.build import mirror_workpaper
from agent.mirror.inventory import build_inventory
from agent.ooxml import ooxml_utils as ox
from agent.schemas import ConclusionOutput, ModelMetadata, SampleItem, SamplePopulationManifest

RED_BOLD = '<a:rPr lang="en-US" sz="1000" b="1"><a:solidFill><a:srgbClr val="FF0000"/></a:solidFill></a:rPr>'
PLAIN = '<a:rPr lang="en-US" sz="1000"/>'


def _anchor(sid: int, name: str, c0: int, r0: int, c1: int, r1: int, body: str, *, txbox: bool, outline: bool = True) -> str:
    ln = '<a:ln w="9525"><a:solidFill><a:srgbClr val="FF0000"/></a:solidFill></a:ln>' if outline else ""
    tx = ' txBox="1"' if txbox else ""
    return (f"<xdr:twoCellAnchor><xdr:from><xdr:col>{c0}</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>{r0}</xdr:row>"
            f"<xdr:rowOff>0</xdr:rowOff></xdr:from><xdr:to><xdr:col>{c1}</xdr:col><xdr:colOff>0</xdr:colOff>"
            f"<xdr:row>{r1}</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:to><xdr:sp macro=\"\" textlink=\"\"><xdr:nvSpPr>"
            f'<xdr:cNvPr id="{sid}" name="{name}"/><xdr:cNvSpPr{tx}/></xdr:nvSpPr>'
            f'<xdr:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="1" cy="1"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/>'
            f"</a:prstGeom><a:noFill/>{ln}</xdr:spPr><xdr:txBody><a:bodyPr vertOverflow=\"clip\" wrap=\"square\"/>"
            f"<a:lstStyle/>{body}</xdr:txBody></xdr:sp><xdr:clientData/></xdr:twoCellAnchor>")


def _p(*runs: tuple[str, str]) -> str:
    if not runs:
        return '<a:p><a:endParaRPr lang="en-US" sz="1000"/></a:p>'
    return "<a:p>" + "".join(f"<a:r>{r}<a:t>{t}</a:t></a:r>" for r, t in runs) + "</a:p>"


def _page(path: Path, lines: list[str]) -> Path:
    im = Image.new("RGB", (900, 500), "white")
    d = ImageDraw.Draw(im)
    try:
        font = ImageFont.truetype("DejaVuSans.ttf", 34)
    except OSError:
        font = ImageFont.load_default()
    for i, line in enumerate(lines):
        d.text((40, 40 + i * 70), line, fill="black", font=font)
    im.save(path)
    return path


def make_py(tmp: Path) -> Path:
    """A PY shaped like CI.9: lead tab + tab '1'."""
    wb = openpyxl.Workbook()
    lead = wb.active
    lead.title = "IA Lead"
    lead["A1"] = " "
    sel = wb.create_sheet("1")
    heads = ["Doc Type", "Doc#", "Invoice Date", "Amount", "Customer Name"]
    sel.append(heads)
    sel.append(["RX", "9578365", "11/05/2025", 17025.40, "Kentucky Exposition Center"])
    for c in "ABCDE":
        sel.column_dimensions[c].width = 18
    py_invoice = _page(tmp / "py_invoice.png", ["INVOICE", "Invoice 9578365", "Total 17,025.40"])
    py_approval = _page(tmp / "py_approval.png", ["APPROVAL", "Approved by Ryan Capito", "on 11/04/2025"])
    sel.add_image(XLImage(str(py_invoice)), "A4")
    sel.add_image(XLImage(str(py_approval)), "H4")
    path = tmp / "PY.xlsx"
    wb.save(path)

    work = tmp / "py_work"
    ox.extract_template(str(path), str(work))
    sheets = ox.resolve_sheets(str(work))
    sel_rels = (work / "xl" / "worksheets" / "_rels" / f"sheet{sheets['1']}.xml.rels").read_text()
    sel_drawing = re.search(r"drawings/(drawing\d+\.xml)", sel_rels).group(1)
    dpath = work / "xl" / "drawings" / sel_drawing
    xml = dpath.read_text()
    root = re.search(r"<((?:\w+:)?wsDr)\b[^>]*>", xml)
    for prefix, uri in (("xdr", "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"),
                        ("a", "http://schemas.openxmlformats.org/drawingml/2006/main")):
        if f"xmlns:{prefix}=" not in root.group(0):
            xml = xml.replace(root.group(0), root.group(0)[:-1] + f' xmlns:{prefix}="{uri}">', 1)
            root = re.search(r"<((?:\w+:)?wsDr)\b[^>]*>", xml)
    shapes = "".join([
        # frames exactly around the two pictures are added below once their extents are known
        _anchor(101, "TextBox 1", 13, 3, 20, 10, _p((RED_BOLD, "Test Step 1: "), (PLAIN, "IA verified that the invoice 9578365 was approved by Ryan Capito on 11/04/2025 prior to payment.")) + _p() + _p((RED_BOLD, "Test Satisfied.")), txbox=True),
        _anchor(102, "TextBox 2", 13, 11, 20, 16, _p((RED_BOLD, "Test Step 2: "), (PLAIN, "IA agreed the invoice total of $17,025.40 to E1. See TM1.")) + _p() + _p((RED_BOLD, "Test Satisfied.")), txbox=True),
        _anchor(103, "TextBox 3", 13, 17, 20, 20, _p((RED_BOLD, "Note: "), (PLAIN, "A credit memo of $679.40 was issued in the prior year.")), txbox=True),
        _anchor(104, "TextBox 4", 5, 8, 6, 9, _p((RED_BOLD, "TM1")), txbox=True),
        _anchor(105, "Rectangle 5", 3, 8, 5, 9, _p(), txbox=False),
    ])
    xml = re.sub(r"(</(?:\w+:)?wsDr>)", lambda m: shapes + m.group(1), xml, count=1)
    dpath.write_text(xml)

    lead_drawing = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                    '<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" '
                    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
                    + _anchor(201, "TextBox 1", 1, 1, 10, 15,
                              _p((RED_BOLD, "Control Description: "), (PLAIN, "Invoices over $5,000 are approved."))
                              + _p() + _p((RED_BOLD, "Test Steps:")) + _p((PLAIN, "1. Verify approval."))
                              + _p((PLAIN, "2. Agree invoice total to E1.")) + _p()
                              + _p((RED_BOLD, "Test Step References:")) + _p((PLAIN, "Refer to Tab '1'.")), txbox=True)
                    + "</xdr:wsDr>")
    (work / "xl" / "drawings" / "drawing90.xml").write_text(lead_drawing)
    rels_dir = work / "xl" / "worksheets" / "_rels"
    (rels_dir / f"sheet{sheets['IA Lead']}.xml.rels").write_text(
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rIdL" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/drawing" '
        'Target="../drawings/drawing90.xml"/></Relationships>')
    lead_sheet = work / "xl" / "worksheets" / f"sheet{sheets['IA Lead']}.xml"
    lx = lead_sheet.read_text()
    if "xmlns:r=" not in lx.split(">", 2)[1]:
        lx = lx.replace("<worksheet ", '<worksheet xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" ', 1)
    lead_sheet.write_text(lx.replace("</worksheet>", '<drawing r:id="rIdL"/></worksheet>'))
    ct = (work / "[Content_Types].xml").read_text()
    (work / "[Content_Types].xml").write_text(ct.replace("</Types>", '<Override PartName="/xl/drawings/drawing90.xml" '
                                                          'ContentType="application/vnd.openxmlformats-officedocument.drawing+xml"/></Types>'))
    # Frames: outlines exactly over each picture.
    inv = build_inventory(str(work))
    tab = next(t for t in inv.tabs if t.name == "1")
    xml = dpath.read_text()
    frames = "".join(
        f'<xdr:absoluteAnchor><xdr:pos x="{p.box.x}" y="{p.box.y}"/><xdr:ext cx="{p.box.cx}" cy="{p.box.cy}"/>'
        f'<xdr:sp macro="" textlink=""><xdr:nvSpPr><xdr:cNvPr id="{300 + i}" name="Rectangle frame {i}"/><xdr:cNvSpPr/></xdr:nvSpPr>'
        f'<xdr:spPr><a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/><a:ln><a:solidFill><a:srgbClr val="FF0000"/>'
        f'</a:solidFill></a:ln></xdr:spPr></xdr:sp><xdr:clientData/></xdr:absoluteAnchor>'
        for i, p in enumerate(tab.by_kind("picture")))
    dpath.write_text(re.sub(r"(</(?:\w+:)?wsDr>)", lambda m: frames + m.group(1), xml, count=1))
    ox.repackage(str(work), str(path))
    return path


def _conclusion(step: str, sample_ids: list[str], result="satisfied") -> ConclusionOutput:
    return ConclusionOutput(
        test_step_id=step, control_objective_ref="CO", conclusion=result, narrative=f"{step} {result}.",
        evidence_citations=[
            {"evidence_id": f"e{i}{s}", "source_file": f, "location": "p.1", "quote_or_summary": q,
             "relevance": "r", "sample_id": s}
            for s in sample_ids
            for i, (f, q) in enumerate([("cy_invoice.png", "Invoice 9850995 total 14,945.60"),
                                        ("cy_approval.png", "Approved by Jane Doe on 06/20/2026")])
        ],
        procedures_performed=["inspection"],
        sample_results=[{"sample_id": s, "conclusion": result} for s in sample_ids],
        attribute_results=[{"attribute": "approval", "sample_id": s, "result": result,
                            "value_observed": "Jane Doe 06/20/2026; total 14,945.60",
                            "evidence_ids": [f"e0{s}", f"e1{s}"]} for s in sample_ids],
        exceptions=[], additional_support_requests=[], confidence="high", confidence_rationale="c",
        sample_coverage=None,
        model_metadata=ModelMetadata(model="m", prompt_version="v1", timestamp="t", tool_call_count=1),
    )


class Stub:
    """Stands in for the model with a fixed, sensible edit: CY facts in, PY
    labels and structure kept, exhibits by role, TM1 on the CY total."""

    def __init__(self, reply=None):
        self.prompts: list[str] = []
        self.reply = reply
        self.messages = self

    def create(self, **kw):
        user = kw["messages"][0]["content"]
        self.prompts.append(user)
        if "reviewing a CY SOX workpaper" in kw["system"]:
            text = '{"flags": []}'
        elif self.reply is not None:
            text = self.reply
        else:
            items = {}
            for m in re.finditer(r"- id (\S+) \(.*?\):\n((?:    \| .*\n?)+)", user):
                items[m.group(1)] = [line[6:] for line in m.group(2).splitlines()]
            text_plan = []
            for iid, paras in items.items():
                new = [p.replace("9578365", "9850995").replace("Ryan Capito", "Jane Doe")
                       .replace("11/04/2025", "06/20/2026").replace("$17,025.40", "$14,945.60") for p in paras]
                if any("credit memo" in p for p in paras):
                    text_plan.append({"id": iid, "action": "remove"})
                else:
                    text_plan.append({"id": iid, "action": "replace", "paragraphs": new})
            pics = re.findall(r"^- id (\S+): (.*)$", user.split("PY PICTURE SLOTS")[1].split("\n\n")[0], re.M) \
                if "PY PICTURE SLOTS" in user else []
            assign = []
            for pid, note in pics:
                assign.append({"id": pid, "exhibit": "X1" if "INVOICE" in note.upper() else "X2"})
            labels = re.findall(r"^- id (\S+): 'TM1'", user, re.M)
            text = json.dumps({
                "text": text_plan, "pictures": assign, "extra_exhibits": [],
                "tickmarks": [{"id": l, "action": "place", "exhibit": "X1", "anchor": "14,945.60"} for l in labels],
                "cells": [], "notes": [],
            })

        class _B:
            type = "text"

        b = _B()
        b.text = text

        class _R:
            content = [b]
            usage = None

        return _R()


def _run(tmp: Path, *, samples=("1",), steps=("TS-1", "TS-2"), stub=None, listing=None):
    py = make_py(tmp)
    _page(tmp / "cy_invoice.png", ["INVOICE", "Invoice 9850995", "Total 14,945.60"])
    _page(tmp / "cy_approval.png", ["APPROVAL", "Approved by Jane Doe", "on 06/20/2026"])
    results = {s: {"conclusion": _conclusion(s, list(samples)), "audit_log": []} for s in steps}
    fields = listing or {"Doc Type": "RX", "Doc#": "9850995", "Invoice Date": "07/01/2026",
                         "Amount": "14945.60", "Customer Name": "Berkowski and Associates Inc"}
    manifests = {"TS-1": SamplePopulationManifest(test_step_id="TS-1", sample_size=len(samples), samples=[
        SampleItem(sample_id=s, test_step_id="TS-1", identifying_details="x", key_fields=dict(fields)) for s in samples])}
    spec = {"control_id": "CI.9", "test_steps": [{"test_step_id": s, "test_step_text": s} for s in steps]}
    stub = stub or Stub()
    out = mirror_workpaper(spec, results, py, tmp / "CY.xlsx", tmp, manifests, stub, "m",
                           summary_rows_fn=None)
    return py, out, stub


# --------------------------------------------------------------------------


def test_the_inventory_reads_the_py_grammar(tmp_path: Path):
    py = make_py(tmp_path)
    work = tmp_path / "w"
    ox.extract_template(str(py), str(work))
    inv = build_inventory(str(work))
    roles = {t.name: t.role for t in inv.tabs}
    assert roles == {"IA Lead": "lead", "1": "selection"}
    tab = next(t for t in inv.tabs if t.name == "1")
    kinds = sorted(i.kind for i in tab.items)
    assert kinds.count("picture") == 2 and kinds.count("frame") == 2
    assert kinds.count("text") == 3 and kinds.count("label") == 1 and kinds.count("box") == 1
    conclusion = next(i for i in tab.items if i.name == "TextBox 1")
    # Formatting kept as style classes: the bold red label is one class, the prose another.
    first = conclusion.paragraphs[0].runs
    assert first[0][1] == "Test Step 1: " and tab.styles[first[0][0]].bold
    assert not tab.styles[first[1][0]].bold


def test_the_cy_file_is_the_py_with_only_the_facts_changed(tmp_path: Path):
    py, out, stub = _run(tmp_path)
    wb = openpyxl.load_workbook(out.path)
    assert wb.sheetnames == ["Mirror Check", "IA Lead", "1"]
    sel = wb["1"]
    # Data row: each CY value under its heading; nothing of last year's left.
    assert [c.value for c in sel[1]][:5] == ["Doc Type", "Doc#", "Invoice Date", "Amount", "Customer Name"]
    # The PY cell decides the kind: the PY held these as text, so CY is text;
    # the PY amount was a number, so CY's is a number (its format applies).
    assert sel["B2"].value == "9850995" and sel["E2"].value == "Berkowski and Associates Inc"
    assert sel["C2"].value == "07/01/2026"
    assert sel["D2"].value == pytest.approx(14945.60)
    wb.close()

    work = tmp_path / "out_w"
    ox.extract_template(str(out.path), str(work))
    assert ox.package_problems(str(work)) == []
    cy = next(t for t in build_inventory(str(work)).tabs if t.name == "1")
    pyw = tmp_path / "py_w2"
    ox.extract_template(str(py), str(pyw))
    pyt = next(t for t in build_inventory(str(pyw)).tabs if t.name == "1")

    # Conclusion boxes: same boxes, same place, same formatting, CY facts.
    for name in ("TextBox 1", "TextBox 2"):
        a = next(i for i in pyt.items if i.name == name)
        b = next(i for i in cy.items if i.name == name)
        assert (a.box.x, a.box.y, a.box.cx, a.box.cy) == (b.box.x, b.box.y, b.box.cx, b.box.cy)
        assert [p.text == "" for p in a.paragraphs] == [p.text == "" for p in b.paragraphs]
        assert b.paragraphs[0].runs[0][1] == a.paragraphs[0].runs[0][1]  # "Test Step N: " label kept
        assert cy.styles[b.paragraphs[0].runs[0][0]].bold
    step1 = next(i for i in cy.items if i.name == "TextBox 1").text
    assert "Jane Doe" in step1 and "06/20/2026" in step1 and "Ryan Capito" not in step1
    # The PY-only note was removed, not left to mislead.
    assert not any(i.name == "TextBox 3" for i in cy.items)

    # Exhibits: CY pages in the PY's slots, each with its frame resized to it.
    pics = cy.by_kind("picture")
    assert len(pics) == 2
    for p in pics:
        slot = next(s for s in pyt.by_kind("picture") if abs(s.box.x - p.box.x) < 10 and abs(s.box.y - p.box.y) < 10)
        assert p.box.cx <= slot.box.cx + 1 and p.box.cy <= slot.box.cy + 1
    assert len(cy.by_kind("frame")) == 2

    # TM1: the PY's own label and box, moved onto the CY total; label beside it.
    label = cy.by_kind("label")[0]
    box = cy.by_kind("box")[0]
    invoice = next(p for p in pics if p.box.x < 2_000_000)
    assert invoice.box.x <= box.box.x and box.box.right <= invoice.box.right + 100_000
    assert not label.box.overlaps(box.box)


def test_one_tab_per_selection_never_per_test_step(tmp_path: Path):
    # Three test steps, one selection: ONE tab. A real output made three.
    _py, out, _s = _run(tmp_path, steps=("TS-1", "TS-2", "TS-3"))
    assert openpyxl.load_workbook(out.path).sheetnames == ["Mirror Check", "IA Lead", "1"]


def test_more_selections_than_py_tabs_copy_the_whole_py_tab(tmp_path: Path):
    _py, out, _s = _run(tmp_path, samples=("1", "2"))
    assert openpyxl.load_workbook(out.path).sheetnames == ["Mirror Check", "IA Lead", "1", "2"]
    work = tmp_path / "w2"
    ox.extract_template(str(out.path), str(work))
    tabs = {t.name: t for t in build_inventory(str(work)).tabs}
    # Tab 2 is a full copy of the PY tab -- conclusion boxes and all -- edited.
    assert {i.name for i in tabs["2"].by_kind("text")} == {"TextBox 1", "TextBox 2"}
    assert ox.package_problems(str(work)) == []


def test_a_failed_model_call_keeps_the_py_text_and_says_so(tmp_path: Path):
    _py, out, _s = _run(tmp_path, stub=Stub(reply="not json at all"))
    assert any("did not return JSON" in n for n in out.notes)
    work = tmp_path / "w3"
    ox.extract_template(str(out.path), str(work))
    assert ox.package_problems(str(work)) == []


def test_py_facts_left_in_cy_text_are_flagged(tmp_path: Path):
    lazy = Stub()
    original = lazy.create

    def keep_py_name(**kw):
        r = original(**kw)
        r.content[0].text = r.content[0].text.replace("Jane Doe", "Ryan Capito").replace("$14,945.60", "$17,025.40")
        return r

    lazy.create = keep_py_name
    _py, out, _s = _run(tmp_path, stub=lazy)
    issues = " ".join(f.issue for f in out.flags)
    assert "Ryan Capito" in issues and "17,025.40" in issues
    sheet = openpyxl.load_workbook(out.path)["Mirror Check"]
    assert any("Ryan Capito" in str(c.value) for row in sheet.iter_rows() for c in row)


def test_a_listing_matching_no_py_heading_leaves_the_py_headings_alone(tmp_path: Path):
    _py, out, _s = _run(tmp_path, listing={"Foo": "1", "Bar": "2"})
    sel = openpyxl.load_workbook(out.path)["1"]
    assert sel["A1"].value == "Doc Type"  # the PY's heading, untouched
    assert any(f.severity == "high" and "none of the PY tab's column headings" in f.issue for f in out.flags)


def test_the_joint_falls_back_when_the_mirror_engine_fails(tmp_path: Path):
    from agent.build_cy_workpaper import build_cy_workpaper

    py = make_py(tmp_path)
    _page(tmp_path / "cy_invoice.png", ["Invoice 9850995"])
    _page(tmp_path / "cy_approval.png", ["Approved"])

    class Broken:
        class messages:
            @staticmethod
            def create(**_kw):
                raise RuntimeError("boom")

    results = {"TS-1": {"conclusion": _conclusion("TS-1", ["1"]), "audit_log": []}}
    spec = {"control_id": "CI.9", "test_steps": [{"test_step_id": "TS-1", "test_step_text": "x"}]}
    outcome = build_cy_workpaper(spec, results, py.name, tmp_path / "o", support_dir=tmp_path,
                                 mirror_client=Broken(), mirror_model="m")
    # A failed call is not a failed build: the PY copy is still made, PY
    # text kept, and the failure reported.
    assert outcome.path.exists()
    assert any("rewrite call failed" in w for w in outcome.warnings) or not outcome.mirrored_py
